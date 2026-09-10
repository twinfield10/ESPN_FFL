"""The Athletic as the sixth blended source.

Modelled on ``test_usage_fifth_source.py``, which is this repo's template for adding
a source: registration, the one real parse bug, abstention, and what must stay out.

The distinguishing property here is that the source is a **file somebody saved**
rather than a scrape, so the tests that matter most are the ones about what happens
when it is absent, stale, or carries a row it should not.

The same workbook also carries Jake Ciely's hand ranking, which is *not* a source: it
projects nothing, casts no vote, and exists to be read beside the board. The section
at the foot of this file is about keeping it that way -- the load-bearing test being
that adding it leaves ``TRUE_Points`` untouched.

**The weekly grain, added 2026-09-09**, is a *second* workbook rather than a second
tab of the first: one sheet of four side-by-side position blocks, nine stats instead
of twelve, and its own abbreviations. Its section is in the middle of this file. The
tests that matter there are the ones about geometry -- the blocks have different
widths, so a parser that trusted column positions would hand tight ends carries.
"""

import numpy as np
import pandas as pd
import pytest

import Scripts.load_athletic as la
import Scripts.projection_utils as pu
import Scripts.refresh_status as rs
import Scripts.season_projections as sp


# --- registration --------------------------------------------------------

def test_ath_is_registered_at_an_equal_vote():
    """0.25, the same as every other external. The equal-vote rule is the invariant;
    `test_usage_fifth_source` holds the general form of it and this pins the value."""
    entry = pu.WEIGHTS["default"]
    assert entry["ATH"] == 0.25
    assert {entry[k] for k in ("ESPN", "FP", "PINNY", "BOL", "ATH")} == {0.25}


def test_ath_is_scored_like_every_other_source():
    import inspect
    default = inspect.signature(pu.proj_to_score).parameters["col_pfix_list"].default
    assert "ATH" in default


def test_ath_counts_toward_coverage_and_toward_the_spread():
    """Both, and for different reasons -- see the two docstrings on those tuples.

    In ``PROJECTION_PREFIXES`` because it moves ``TRUE_Points``, so a player only it
    projects must not read as unprojected. In ``OPINION_PREFIXES`` because it answers
    the same question the other externals do, which is the membership rule there.
    """
    assert "ATH" in sp.PROJECTION_PREFIXES
    assert "ATH" in sp.OPINION_PREFIXES


def test_ath_is_withdrawn_on_availability_like_the_other_player_keyed_sources():
    assert "ATH_" in sp.AVAILABILITY_WITHDRAWN_PREFIXES


def test_the_source_is_named_in_the_freshness_manifest():
    """It has no nightly stage, which is *why* this matters: nothing else would ever
    notice the file going stale, and it carries a sixth of every projection it covers.
    """
    named = {name for name, _, _ in rs.PROJECTION_SOURCES}
    assert "The Athletic" in named
    resolve = next(r for n, r, _ in rs.PROJECTION_SOURCES if n == "The Athletic")
    path = resolve(2999)
    assert "2999" in str(path) and path.suffix == ".parquet"
    assert not path.parent.exists(), "resolving a path must not create a directory"


def test_ath_has_a_column_on_the_board():
    """A source that moves `Us` must be readable beside it, and before it."""
    import sys
    sys.path.insert(0, "app")
    import draft_view as dv

    labels = [c.label for c in dv.COLUMNS if c.group == "Points"]
    assert "ATH" in labels
    assert labels.index("ATH") < labels.index("Us")
    spec = next(c for c in dv.COLUMNS if c.source == "ATH_Points")
    assert spec.positions == (), "it covers four positions, so it is not scoped"
    assert spec.source_of and spec.how and spec.caveat


# --- the parse -----------------------------------------------------------

def test_the_stat_map_uses_espn_names_matching_fantasypros():
    """Both sources must land on identical ``<PREFIX>_<stat>`` columns or the blend
    is comparing two different vocabularies."""
    assert set(la.STAT_COLUMNS.values()) == {
        "passingAttempts", "passingCompletions", "passingYards",
        "passingTouchdowns", "passingInterceptions",
        "rushingAttempts", "rushingYards", "rushingTouchdowns",
        "receivingTargets", "receivingReceptions", "receivingYards",
        "receivingTouchdowns",
    }
    assert "lostFumbles" not in la.STAT_COLUMNS.values(), (
        "the workbook does not project fumbles; ATH_ must abstain rather than "
        "carry a column of zeroes")


def test_all_thirty_two_team_tabs_are_read():
    assert len(la.TEAM_TABS) == 32
    assert len(set(la.TEAM_TABS)) == 32
    # ESPN's abbreviations, not nflverse's -- so no alias map is needed on this path.
    assert {"WSH", "JAX", "LV", "LAR", "LAC"} <= set(la.TEAM_TABS)
    assert "WAS" not in la.TEAM_TABS and "LA" not in la.TEAM_TABS


def test_a_quarterback_cannot_carry_receiving_stats():
    """The one real bug in the file, as an executable fact.

    The workbook splits team target share across a tab's rows and on the New Orleans
    tab some of it lands on Spencer Rattler, a third-string quarterback: 32.2 targets
    and 258.7 receiving yards. Read straight he scores 59.8 instead of 7.8, which
    would make him a real opinion in the blend.
    """
    assert "receivingYards" not in la.POSITION_STATS["QB"]
    assert "receivingReceptions" not in la.POSITION_STATS["QB"]
    assert "receivingTargets" not in la.POSITION_STATS["QB"]
    # And the converse: no skill position may carry passing stats.
    for pos in ("RB", "WR", "TE"):
        assert not (la.POSITION_STATS[pos] & {
            "passingYards", "passingTouchdowns", "passingAttempts"})


def test_tight_ends_are_not_given_carries():
    """Not defensive coding -- the same share model that produced the Rattler row
    could allocate rush share to a tight end on the next download."""
    assert "rushingAttempts" not in la.POSITION_STATS["TE"]


def test_defence_and_kicker_are_not_ingested():
    """The workbook's DST tab defines all seven points-allowed tiers in `Settings`
    and leaves the bucket columns null for all 32 teams, so its own defence values
    silently omit that component. This repo's DST model is blended at 0.25 instead."""
    assert set(la.POSITION_STATS) == {"QB", "RB", "WR", "TE"}
    assert "D/ST" not in la.POSITION_STATS and "K" not in la.POSITION_STATS


def test_the_masked_stats_flag_stays_out_of_the_blend_namespace():
    """``UPPER_`` is reserved for blendable numerics -- ``compute_weighted_stats`` and
    ``proj_to_score`` scan every uppercase prefix and require it to be numeric."""
    for col in la.DIAGNOSTIC_COLUMNS:
        assert col.islower(), col
        assert not col.startswith(la.PREFIX)


# --- absence and abstention ---------------------------------------------

def test_a_missing_workbook_degrades_rather_than_raising(tmp_path, monkeypatch,
                                                         capsys):
    """Every other loader returns an empty frame with a fix hint when its file is
    absent, and ``build_season_projections`` skips it. A hand-dropped file is the one
    most likely to be missing, so this is the path that must not raise."""
    monkeypatch.setattr(sp, "season_dir",
                        lambda *a, **k: tmp_path / "nope.parquet")
    out = sp.load_theathletic_season(2026)
    assert out.empty and list(out.columns) == ["name_key"]
    assert "load_athletic" in capsys.readouterr().out


def test_an_unmatched_player_abstains_rather_than_projecting_zero():
    """``ATH_`` is in the imputation chain so its gaps arrive flagged. Without a
    provenance column an unmatched row enters as a confident projection of zero and
    drags the player toward it -- the trap ``test_usage_fifth_source`` names.
    """
    frame = pd.DataFrame({
        "ESPN_rushingYards": [100.0],
        "FP_rushingYards": [120.0],
        "ATH_rushingYards": [np.nan],
    })
    frame = pu.impute_columns(frame, target_prefix="ATH_", source_prefix="MEAN_")
    if "ATH_rushingYards_is_imputed" not in frame.columns:
        frame["ATH_rushingYards_is_imputed"] = True
    weights = {"default": {"ESPN": 0.5, "FP": 0.5, "ATH": 0.5}}
    out = pu.compute_weighted_stats(frame, ["rushingYards"], weights)
    assert out["TRUE_rushingYards"][0] == pytest.approx(110.0)


def test_a_real_line_gets_its_equal_vote():
    frame = pd.DataFrame({
        "ESPN_rushingYards": [100.0],
        "FP_rushingYards": [100.0],
        "ATH_rushingYards": [400.0],
        "ESPN_rushingYards_is_imputed": [False],
        "FP_rushingYards_is_imputed": [False],
        "ATH_rushingYards_is_imputed": [False],
    })
    weights = {"default": {"ESPN": 0.25, "FP": 0.25, "ATH": 0.25}}
    out = pu.compute_weighted_stats(frame, ["rushingYards"], weights)
    assert out["TRUE_rushingYards"][0] == pytest.approx(200.0)


# --- the weekly grain ----------------------------------------------------
#
# A second workbook with its own geometry, and the registration is separate from the
# season one above: `WEIGHTS` was already shared, so adding `ATH` to `WEEKLY_PREFIXES`
# is what actually turned the fifth weekly vote on. See docs/plans/47-athletic-weekly.md.


def _weekly_sheet_rows(blocks=None, week=1, season=2026):
    """A synthetic weekly sheet with the real geometry.

    The geometry *is* the thing under test, so this fixture reproduces it rather than
    simplifying it: a sparse banner row, a header row, blank spacer columns between
    blocks, and **blocks of different widths** -- the receiver block carries no
    ``Rush Att`` and the tight-end block no rushing columns at all, exactly as the
    shipped file does.

    Args:
        blocks: ``[(position, [headers], [[values], ...]), ...]``. Defaults to the
            four-block shape of the 2026 week-1 download.
        week: Week for the sheet name.
        season: Season for the sheet name.

    Returns:
        tuple: ``(sheet_name, rows)`` ready to hand to a fake workbook.
    """
    if blocks is None:
        blocks = [
            ("QB", ["Name", "Team", "Opp", "Pass YD", "Pass TD", "INT",
                    "Rush Att", "Rush YD", "Rush TD", "FPS"],
             [["Joe Burrow", "CIN", "vs TB", 267.2, 2.3, 0.8, 2.5, 7.7, 0.1, 19.3]]),
            ("RB", ["Name", "Team", "Opp", "Rush Att", "Rush YD", "Rush TD",
                    "REC", "REC YD", "REC TD", "FPS"],
             [["Jahmyr Gibbs", "DET", "vs NO", 17.3, 90.7, 0.9, 4.6, 30.7, 0.2, 21.4],
              ["Bijan Robinson", "ATL", "@ PIT", 19.1, 91.7, 0.6, 4.3, 32.8, 0.2, 19.4]]),
            ("WR", ["Name", "Team", "Opp", "Rush YD", "Rush TD",
                    "REC", "REC YD", "REC TD", "FPS"],
             [["Ja'Marr Chase", "CIN", "vs TB", 1, 0, 6.6, 85.7, 0.7, 16.4],
              ["Puka Nacua", "LAR", "vs SF", 2.9, 0, 7.3, 96.1, 0.5, 16.8],
              ["Malik Benson", "LV", "vs MIA", 0, 0, 0.4, 4.4, 0, 0.8]]),
            ("TE", ["Name", "Team", "Opp", "REC", "REC YD", "REC TD", "FPS"],
             [["Colston Loveland", "CHI", "@ CAR", 5.1, 62.8, 0.5, 11.6]]),
        ]

    width = sum(len(headers) + 1 for _, headers, _ in blocks) - 1
    depth = max(len(players) for _, _, players in blocks)
    banner = [None] * width
    header = [None] * width
    body = [[None] * width for _ in range(depth)]

    at = 0
    for position, headers, players in blocks:
        banner[at] = position
        for offset, head in enumerate(headers):
            header[at + offset] = head
        for row, values in enumerate(players):
            for offset, value in enumerate(values):
                body[row][at + offset] = value
        at += len(headers) + 1                     # the blank spacer column

    name = f"NFL_{season}_Week_{week}_Half_PPR_Weekly"
    return name, [tuple(banner), tuple(header), *(tuple(r) for r in body)]


class _FakeSheet:
    def __init__(self, rows):
        self._rows = rows

    def iter_rows(self, values_only=True):
        return iter(self._rows)


class _FakeBook:
    """Stands in for an ``openpyxl`` workbook, so no .xlsx has to be written."""

    def __init__(self, sheets):
        self._sheets = sheets

    @property
    def sheetnames(self):
        return list(self._sheets)

    def __getitem__(self, name):
        return _FakeSheet(self._sheets[name])

    def close(self):
        pass


@pytest.fixture
def fake_weekly(monkeypatch, tmp_path):
    """Patch ``openpyxl.load_workbook`` to serve a synthetic weekly sheet."""
    import openpyxl

    def install(sheets, filename="Week_1_Proj_0909.xlsx"):
        path = tmp_path / filename
        path.write_bytes(b"not really xlsx")
        monkeypatch.setattr(openpyxl, "load_workbook",
                            lambda *a, **k: _FakeBook(sheets))
        return path

    return install


def test_the_weekly_source_is_registered_as_the_fifth_weekly_vote():
    """One line of data, because `WEIGHTS` is shared between the grains."""
    assert "ATH" in pu.WEEKLY_PREFIXES
    assert "theathletic" in pu.WEEKLY_SOURCE_FILES
    assert pu.WEIGHTS["default"]["ATH"] == 0.25
    # `proj_to_score` already listed ATH for the season path, so nothing there moved.
    import inspect
    default = inspect.signature(pu.proj_to_score).parameters["col_pfix_list"].default
    assert "ATH" in default


def test_the_weekly_source_is_shown_rather_than_silently_blended():
    """A source with an equal vote and no column reads as agreement -- this repo's
    oldest failure mode. `real_sources` is what the weekly table's column list and
    the absent-source caption both come from."""
    from app import lineup as lu
    from app import lineup_table as ltab

    assert ("ATH", "theathletic") in lu.WEEKLY_SOURCES
    assert "ATH" in ltab.SOURCE_HELP and "ATH" in ltab.POINTS_LABELS
    absent = {"weekly_sources_present": {"theathletic": False}}
    assert "ATH" not in lu.real_sources(absent)
    assert "ATH" in lu.real_sources({"weekly_sources_present": {"theathletic": True}})


def test_the_weekly_source_is_named_in_the_freshness_manifest():
    """A hand download is the one that goes missing, so it must be watched."""
    entry = next(e for e in rs.WEEKLY_PROJECTION_SOURCES if e[0] == "ATH weekly")
    assert "load_athletic" in entry[2] and "--what weekly" in entry[2]
    assert entry[3] is False, "an equal vote going missing is a real fault"


def test_the_blocks_are_read_by_header_not_by_position(fake_weekly):
    """The load-bearing parse test. The blocks have different widths -- 10 columns for
    the backs, 9 for the receivers, 7 for the tight ends -- so a parser that trusted
    positions would read a receiver's ``REC`` into a tight end's ``Rush Att``."""
    name, rows = _weekly_sheet_rows()
    path = fake_weekly({name: rows})
    frame, season, week, flavor = la.read_weekly_workbook(path)

    assert (season, week, flavor) == (2026, 1, "Half_PPR")
    assert len(frame) == 7
    assert dict(frame["position"].value_counts()) == {"WR": 3, "RB": 2, "QB": 1, "TE": 1}

    burrow = frame[frame["player_name"] == "Joe Burrow"].iloc[0]
    assert burrow["proj_passingYards"] == 267.2
    assert burrow["proj_rushingAttempts"] == 2.5
    gibbs = frame[frame["player_name"] == "Jahmyr Gibbs"].iloc[0]
    assert gibbs["proj_receivingReceptions"] == 4.6
    assert gibbs["proj_rushingAttempts"] == 17.3


def test_a_receiver_gets_no_carries_and_a_tight_end_no_rushing_at_all(fake_weekly):
    """Not the position mask -- the *workbook* has no such column for them. Reading
    one anyway is what a positional parser does."""
    name, rows = _weekly_sheet_rows()
    path = fake_weekly({name: rows})
    frame, _, _, _ = la.read_weekly_workbook(path)

    wr = frame[frame["position"] == "WR"]
    assert wr["proj_rushingAttempts"].isna().all(), "the WR block has no Rush Att"
    assert wr["proj_rushingYards"].notna().all(), "but it does have Rush YD"

    te = frame[frame["position"] == "TE"]
    for stat in ("proj_rushingAttempts", "proj_rushingYards", "proj_rushingTouchdowns"):
        assert te[stat].isna().all(), f"the TE block has no {stat}"
    assert te["proj_receivingReceptions"].notna().all()


def test_the_position_mask_still_fires_if_the_workbook_drifts(fake_weekly):
    """The mask is near-trivially satisfied by the current geometry, which is exactly
    why it must be tested against geometry that would defeat it: a tight-end block
    that gains a rush column."""
    name, rows = _weekly_sheet_rows(blocks=[
        ("TE", ["Name", "Team", "Opp", "Rush Att", "Rush YD", "REC", "REC YD"],
         [["Rogue Tight End", "CHI", "@ CAR", 9.0, 44.0, 5.1, 62.8]]),
    ])
    path = fake_weekly({name: rows})
    frame, _, _, _ = la.read_weekly_workbook(path)

    row = frame.iloc[0]
    assert pd.isna(row["proj_rushingAttempts"]) and pd.isna(row["proj_rushingYards"])
    assert row["proj_receivingReceptions"] == 5.1
    assert set(row["masked_stats"].split(",")) == {"rushingAttempts", "rushingYards"}


def test_the_workbook_s_own_points_are_never_read(fake_weekly):
    """``FPS`` is half-PPR and derived from the nine columns beside it. Points are
    what a league's rules do to a stat line -- the rule this module opens with."""
    name, rows = _weekly_sheet_rows()
    path = fake_weekly({name: rows})
    frame, _, _, _ = la.read_weekly_workbook(path)

    assert "FPS" not in la.WEEKLY_STAT_COLUMNS
    assert not [c for c in frame.columns
                if "FPS" in c or c.endswith("_Points") or c == "Points"]


def test_the_weekly_stats_are_the_nine_the_workbook_publishes():
    """Three fewer than the season book, and the omissions are load-bearing: with no
    ``receivingTargets`` column the weekly ``MEAN_receivingTargets`` stays ESPN
    alone, which is the gap the season registration was made to close."""
    assert set(la.WEEKLY_STAT_COLUMNS.values()) == {
        "passingYards", "passingTouchdowns", "passingInterceptions",
        "rushingAttempts", "rushingYards", "rushingTouchdowns",
        "receivingReceptions", "receivingYards", "receivingTouchdowns",
    }
    missing = set(la.STAT_COLUMNS.values()) - set(la.WEEKLY_STAT_COLUMNS.values())
    assert missing == {"passingAttempts", "passingCompletions", "receivingTargets"}


def test_the_weekly_teams_are_normalised_to_espn(fake_weekly):
    """The two workbooks disagree with each other: the weekly sheet says ``JAC`` and
    ``WAS`` where the season book's team tabs say ``JAX`` and ``WSH``. Left alone the
    bye-week check reports Jacksonville absent from a week it played."""
    name, rows = _weekly_sheet_rows(blocks=[
        ("TE", ["Name", "Team", "Opp", "REC", "REC YD", "REC TD"],
         [["Brenton Strange", "JAC", "vs CAR", 3.4, 40.0, 0.3],
          ["Zach Ertz", "WAS", "@ GB", 3.0, 30.0, 0.2],
          ["George Kittle", "SF", "vs LAR", 4.0, 50.0, 0.3]]),
    ])
    path = fake_weekly({name: rows})
    frame, _, _, _ = la.read_weekly_workbook(path)
    assert set(frame["pro_team"]) == {"JAX", "WSH", "SF"}


def test_the_sheet_name_is_the_authority_on_the_week(fake_weekly):
    """A hand download does not fail by failing to parse, it fails by being last
    week's copy. The filename is a human's label; the sheet name is the publisher's."""
    name, rows = _weekly_sheet_rows(week=4)
    path = fake_weekly({name: rows}, filename="Week_1_Proj_0909.xlsx")
    frame, _, week, _ = la.read_weekly_workbook(path)
    assert week == 4 and set(frame["week"]) == {4}


def _mk(root, name):
    """A writable output path under ``root``, parents created.

    Stands in for both ``season_dir`` and ``landing_dir`` so a build writes into
    ``tmp_path`` instead of ``Data/Projections``.
    """
    out = root / "out"
    out.mkdir(parents=True, exist_ok=True)
    return out / name


def test_a_week_that_disagrees_with_the_sheet_raises_unless_forced(fake_weekly,
                                                                  monkeypatch,
                                                                  tmp_path, capsys):
    name, rows = _weekly_sheet_rows(week=1)
    path = fake_weekly({name: rows})
    monkeypatch.setattr(la, "season_dir", lambda *a, **k: _mk(tmp_path, a[-1]))
    monkeypatch.setattr(la, "landing_dir", lambda *a, **k: _mk(tmp_path, a[-1]))

    with pytest.raises(ValueError, match="week 1"):
        la.build_weekly(path, week=2)

    written = la.build_weekly(path, week=2, force=True)
    assert set(written["week"]) == {2}
    assert "--force" in capsys.readouterr().out


def test_a_sheet_that_is_not_a_weekly_slate_is_refused(fake_weekly):
    """Named rather than guessed at: the season workbook has 35 tabs and none of them
    is a slate, so handing it to the weekly importer must say so."""
    _, rows = _weekly_sheet_rows()
    path = fake_weekly({"ARI": rows, "Rankings": rows, "Settings": rows})
    with pytest.raises(KeyError, match="exactly one sheet"):
        la.read_weekly_workbook(path)


def test_two_slates_in_one_workbook_are_refused(fake_weekly):
    """Choosing which week to import is not a choice the parser should make."""
    n1, r1 = _weekly_sheet_rows(week=1)
    n2, r2 = _weekly_sheet_rows(week=2)
    path = fake_weekly({n1: r1, n2: r2})
    with pytest.raises(KeyError, match="found 2"):
        la.read_weekly_workbook(path)


def test_the_newest_download_wins_for_its_own_week_only(fake_weekly, monkeypatch,
                                                        tmp_path):
    """The owner's rule, and the departure from `scrape_FP`'s `keep="first"` freeze:
    a re-download mid-week is the normal case for a hand-dropped file. What must not
    happen is it touching another week."""
    monkeypatch.setattr(la, "season_dir", lambda *a, **k: _mk(tmp_path, a[-1]))
    monkeypatch.setattr(la, "landing_dir", lambda *a, **k: _mk(tmp_path, a[-1]))

    n1, r1 = _weekly_sheet_rows(week=1)
    la.build_weekly(fake_weekly({n1: r1}))
    n2, r2 = _weekly_sheet_rows(week=2)
    la.build_weekly(fake_weekly({n2: r2}, filename="Week_2_Proj_0916.xlsx"))

    # Week 1 again, with a different number in it.
    _, revised = _weekly_sheet_rows(week=1, blocks=[
        ("QB", ["Name", "Team", "Opp", "Pass YD", "Pass TD", "INT",
                "Rush Att", "Rush YD", "Rush TD", "FPS"],
         [["Joe Burrow", "CIN", "vs TB", 999.0, 9.9, 0.0, 0.0, 0.0, 0.0, 99.9]]),
    ])
    out = la.build_weekly(fake_weekly({n1: revised},
                                      filename="Week_1_Proj_0912.xlsx"))

    assert sorted(out["week"].unique()) == [1, 2]
    wk1 = out[out["week"] == 1]
    assert len(wk1) == 1 and wk1.iloc[0]["proj_passingYards"] == 999.0
    assert len(out[out["week"] == 2]) == 7, "week 2 must be untouched"


def test_the_file_stays_cumulative(fake_weekly, monkeypatch, tmp_path):
    """`clean_lineups` re-merges this file onto every week in the lineup frame, so a
    current-week-only file would blank The Athletic for prior weeks and turn stored
    history into a four-source board retroactively."""
    monkeypatch.setattr(la, "season_dir", lambda *a, **k: _mk(tmp_path, a[-1]))
    monkeypatch.setattr(la, "landing_dir", lambda *a, **k: _mk(tmp_path, a[-1]))

    for week in (1, 2, 3):
        name, rows = _weekly_sheet_rows(week=week)
        out = la.build_weekly(fake_weekly({name: rows},
                                          filename=f"Week_{week}_Proj.xlsx"))
    assert sorted(out["week"].unique()) == [1, 2, 3]

    # And --no-merge is the deliberate escape hatch, not the default.
    name, rows = _weekly_sheet_rows(week=4)
    out = la.build_weekly(fake_weekly({name: rows}, filename="Week_4_Proj.xlsx"),
                          merge=False)
    assert sorted(out["week"].unique()) == [4]


def test_an_absent_weekly_file_degrades_rather_than_raising(tmp_path, monkeypatch):
    """The pre-season and pre-first-download state, and the state every Tuesday
    before the workbook is saved."""
    monkeypatch.setattr(pu, "theathletic_weekly_parquet",
                        lambda season: tmp_path / "nope.parquet")
    with pytest.warns(pu.MissingProjectionSourceWarning, match="The Athletic"):
        out = pu.clean_ath_weekly(season=2026)
    assert out.empty
    assert list(out.columns) == pu.SOURCE_JOIN_KEYS
    assert out["week"].dtype == "int64", "week is a merge key on both sides"


def test_a_hand_download_is_not_judged_by_the_nightly_staleness_window():
    """It publishes weekly, so a workbook imported on Wednesday is *correctly* four
    days old on Sunday. Under the 48-hour window it would warn every weekend, and a
    warning that fires on the normal case is one nobody reads."""
    assert pu.MANUAL_STALE_AFTER_HOURS > 7 * 24
    assert pu.MANUAL_STALE_AFTER_HOURS > pu.STALE_AFTER_HOURS
    import inspect
    body = inspect.getsource(pu.clean_ath_weekly)
    assert "MANUAL_STALE_AFTER_HOURS" in body


def test_the_weekly_loader_keeps_lowercase_diagnostics_out_of_the_blend(tmp_path,
                                                                       monkeypatch):
    """``UPPER_`` is the blendable namespace: `compute_weighted_stats` and
    `proj_to_score` scan every uppercase prefix and require it to be numeric. A
    ``masked_stats`` string riding in as ``ATH_masked_stats`` would break both."""
    path = tmp_path / "weekly.parquet"
    pd.DataFrame({
        "week": pd.Series([1], dtype="int64"),
        "player_name": ["Joe Burrow"],
        "pro_team": ["CIN"],
        "position": ["QB"],
        "masked_stats": ["receivingYards"],
        "proj_passingYards": [267.2],
    }).to_parquet(path)

    out = pu.clean_ath_weekly(ath_path=path)
    assert list(out.columns) == ["week", "player_name", "proj_passingYards"]
    renamed = pu.change_col_prefix(out, old_pfix="proj", new_pfix="ATH")
    assert [c for c in renamed.columns if c.startswith("ATH_")] == ["ATH_passingYards"]


def test_every_weekly_source_including_this_one_is_aligned_before_it_is_merged():
    """Extends `tests/test_name_audit.py`'s guard to the fifth source.

    The weekly merges key on the raw ``player_name`` string, so ``Kyle Pitts`` against
    ESPN's ``Kyle Pitts Sr.`` is a miss rather than a near miss -- the player abstains
    and the board shows a source agreeing with ESPN about someone it never projected.
    """
    import inspect
    body = inspect.getsource(pu.clean_lineups)
    aligned = body.index("align_to_espn_names(ath_proj")
    merged = body.index("mean_df.merge(ath_proj")
    assert aligned < merged, "ath_proj is merged before it is aligned"


def test_the_weekly_grain_makes_the_owed_measurement_runnable():
    """What plan 38 recorded as impossible. `Scripts.usage.evalset.SOURCES` scores
    player-week rows out of `lineups.parquet`, so a season-only source could never
    appear in it -- and reading a clean `Scripts.lab.accuracy` table would have looked
    like The Athletic had been judged when it had not."""
    from Scripts.usage import evalset as es
    assert "ATH" in es.SOURCES
    assert "ATH" in es.CARRIED


# --- the real weekly file, when it is there ------------------------------

def test_the_shipped_weekly_file_parses_to_offence_only():
    """Skipped where absent -- it is a manual download. Where present, these are the
    numbers the 2026 week-1 import printed."""
    from Scripts.paths import season_dir

    path = season_dir("TheAthletic", 2026, la.WEEKLY_FILENAME, create=False)
    if not path.exists():
        pytest.skip("no 2026 Athletic weekly workbook imported")

    df = pd.read_parquet(path)
    assert set(df["position"]) <= {"QB", "RB", "WR", "TE"}, "no kickers, no defences"
    assert df.groupby("week")["player_name"].apply(lambda s: s.is_unique).all()
    assert df["week"].dtype == "int64"

    qb = df[df["position"] == "QB"]
    assert qb["proj_receivingReceptions"].isna().all(), (
        "a quarterback with receptions means the position mask regressed")
    te = df[df["position"] == "TE"]
    assert te["proj_rushingYards"].isna().all(), (
        "the tight-end block publishes no rushing columns")


# --- the real file, when it is there -------------------------------------

def test_the_shipped_file_parses_to_offence_only():
    """Skipped where the file is absent -- it is a manual download, so CI may not
    have one. Where it is present, these are the numbers the import printed."""
    from Scripts.paths import season_dir

    path = season_dir("TheAthletic", 2026, la.FILENAME, create=False)
    if not path.exists():
        pytest.skip("no 2026 Athletic workbook imported")

    df = pd.read_parquet(path)
    assert set(df["position"]) == {"QB", "RB", "WR", "TE"}
    assert df["player_name"].is_unique
    qb = df[df["position"] == "QB"]
    assert qb["receivingYards"].isna().all(), (
        "a quarterback with receiving yards means the position mask regressed")
    assert df["passingYards"].notna().sum() == len(qb)


# --- the hand ranking, which is not a source -----------------------------

def _ranked_frame():
    """A scored frame with two positions, three of whom Jake ranked."""
    return pd.DataFrame({
        "player_name": ["A", "B", "C", "D", "E"],
        "primaryPosition": ["RB", "RB", "RB", "WR", "WR"],
        "TRUE_Points": [300.0, 200.0, 100.0, 250.0, 150.0],
        "TRUE_PosRank": [1.0, 2.0, 3.0, 1.0, 2.0],
        "ATH_Points": [300.0, 200.0, 100.0, 250.0, 150.0],
        # He agrees on A, has C above B, and never ranked D or E.
        "ath_pos_rank": [1.0, 3.0, 2.0, np.nan, np.nan],
    })


def test_the_rank_columns_are_lowercase_so_the_blend_cannot_see_them():
    """The load-bearing naming rule. ``compute_weighted_stats`` and ``proj_to_score``
    scan every uppercase ``<PREFIX>_`` and require it to be numeric -- and a rank *is*
    numeric, so an uppercase name would get it blended into a stat line and priced.
    Lowercase makes "this never moves TRUE_Points" true by construction."""
    from Scripts.load_athletic import RANK_COLUMNS

    for column in (*RANK_COLUMNS.values(), "ath_pos_rank", "ath_override",
                   "ath_rank_delta"):
        assert column == column.lower(), f"{column} would enter the blend namespace"


def test_the_ranking_is_not_registered_as_a_source():
    """It projects nothing, so it cannot vote, cannot widen the spread, and has
    nothing to withdraw on an injury."""
    assert "ath_pos_rank" not in pu.WEIGHTS["default"]
    for tup in (sp.OPINION_PREFIXES, sp.PROJECTION_PREFIXES,
                sp.AVAILABILITY_WITHDRAWN_PREFIXES):
        assert not any(str(prefix).startswith("ath_") for prefix in tup)


def test_adding_the_rank_does_not_move_a_single_blended_point():
    """The whole basis for shipping this in draft week, asserted rather than promised.

    Run the blend and the scoring over a frame with and without the rank column and
    require the outputs to be identical -- not close, identical.
    """
    base = pd.DataFrame({
        "primaryPosition": ["RB", "WR"],
        "ESPN_rushingYards": [1000.0, 40.0],
        "FP_rushingYards": [1100.0, 60.0],
    })
    weights = {"default": {"ESPN": 0.25, "FP": 0.25}}
    plain = pu.compute_weighted_stats(df=base.copy(), stats_list=["rushingYards"],
                                      weights_dict=weights)
    with_rank = base.copy()
    with_rank["ath_pos_rank"] = [1.0, 12.0]
    ranked = pu.compute_weighted_stats(df=with_rank, stats_list=["rushingYards"],
                                       weights_dict=weights)
    pd.testing.assert_series_equal(plain["TRUE_rushingYards"],
                                   ranked["TRUE_rushingYards"])
    assert "TRUE_ath_pos_rank" not in ranked.columns
    assert "ath_pos_rank_Points" not in ranked.columns


def test_the_override_is_his_rank_against_his_own_projection():
    """Positive means he ranks a player above his own numbers -- the direction
    ``USG_PosRankDelta`` already set for a named voice."""
    out = sp._attach_athletic_override(_ranked_frame())
    # C is his RB2 and his projection's RB3, so he is one spot higher on him.
    assert out.loc[2, "ath_override"] == pytest.approx(1.0)
    assert out.loc[1, "ath_override"] == pytest.approx(-1.0)
    assert out.loc[0, "ath_override"] == pytest.approx(0.0)


def test_a_player_he_never_ranked_gets_no_delta_rather_than_a_number():
    """The receivers are unranked. A zero there would read as agreement."""
    out = sp._attach_athletic_override(_ranked_frame())
    assert out.loc[[3, 4], "ath_override"].isna().all()


def test_the_override_ranks_him_only_against_the_players_he_ranked():
    """He ranks 85 backs and a board carries half as many again. Ranking ``ATH_Points``
    over the whole pool would score his 60th back against a 60th drawn from a deeper
    one and report bench depth as disagreement."""
    frame = _ranked_frame()
    # Two more backs he never looked at, both projected above everyone he did.
    deeper = pd.concat([frame, pd.DataFrame({
        "player_name": ["X", "Y"], "primaryPosition": ["RB", "RB"],
        "TRUE_Points": [400.0, 350.0], "TRUE_PosRank": [1.0, 2.0],
        "ATH_Points": [400.0, 350.0], "ath_pos_rank": [np.nan, np.nan],
    })], ignore_index=True)
    out = sp._attach_athletic_override(deeper)
    # Unchanged: the two interlopers cannot push his ranked backs down.
    assert out.loc[2, "ath_override"] == pytest.approx(1.0)
    assert out.loc[1, "ath_override"] == pytest.approx(-1.0)


def test_a_frame_with_no_ranking_keeps_its_shape():
    """A board built before the ranking existed carries no column, and asking for the
    override must not invent one."""
    frame = _ranked_frame().drop(columns=["ath_pos_rank"])
    out = sp._attach_athletic_override(frame)
    assert "ath_override" not in out.columns


@pytest.mark.parametrize("points,flavor", [
    (0.0, "std"), (0.5, "half"), (1.0, "ppr"),
    # Between two lists, the nearer one. A 0.75-PPR league is not a case the workbook
    # has a list for, and the nearest list beats no list.
    # Exactly between two lists, the lower one -- see `_athletic_rank_flavor`.
    (0.75, "half"), (0.2, "std"),
])
def test_the_flavor_follows_this_league_s_points_per_reception(
        points, flavor, monkeypatch):
    """The workbook ranks the same 290 players three times because a reception is
    worth a different amount in each. Reading the one that answers this league's rules
    is the same principle that scores every source through them."""
    monkeypatch.setattr(sp, "_modelled_scoring_weights",
                        lambda league, season: {"receivingReceptions": points})
    assert sp._athletic_rank_flavor(object(), 2026) == flavor


# --- the real ranking, when it is there ----------------------------------

def test_the_shipped_ranking_covers_the_four_offensive_positions():
    from Scripts.paths import season_dir

    path = season_dir("TheAthletic", 2026, la.RANKS_FILENAME, create=False)
    if not path.exists():
        pytest.skip("no 2026 Athletic workbook imported")

    ranks = pd.read_parquet(path)
    assert ranks["player_name"].is_unique
    assert ranks.groupby("position").size().to_dict() == {
        "QB": 40, "RB": 85, "TE": 45, "WR": 120}
    # Each list is a dense 1..N with no ties, per position.
    for column in la.RANK_COLUMNS.values():
        for position, block in ranks.groupby("position"):
            assert sorted(block[column]) == list(range(1, len(block) + 1)), (
                f"{column} is not a dense ranking of the {position}s")


def test_every_ranked_player_is_also_a_projected_player():
    """Both files come out of one ``build`` call off one workbook, so a name on the
    ranking that the team tabs do not carry means the workbook spells him two ways --
    a join miss that would cost a board column."""
    from Scripts.paths import season_dir

    ranks_path = season_dir("TheAthletic", 2026, la.RANKS_FILENAME, create=False)
    stats_path = season_dir("TheAthletic", 2026, la.FILENAME, create=False)
    if not (ranks_path.exists() and stats_path.exists()):
        pytest.skip("no 2026 Athletic workbook imported")

    ranks = pd.read_parquet(ranks_path)
    stats = pd.read_parquet(stats_path)
    orphans = set(ranks["player_name"]) - set(stats["player_name"])
    assert not orphans, f"ranked but not projected: {sorted(orphans)}"


def test_he_overrides_his_own_numbers_where_projections_are_weakest():
    """The shape is the argument for showing the column at all: he leaves the position
    a projection handles best alone and reworks the two it handles worst. Measured on
    the 2026-08-31 workbook, in his own half-PPR list.
    """
    from Scripts.paths import season_dir

    ranks_path = season_dir("TheAthletic", 2026, la.RANKS_FILENAME, create=False)
    stats_path = season_dir("TheAthletic", 2026, la.FILENAME, create=False)
    if not (ranks_path.exists() and stats_path.exists()):
        pytest.skip("no 2026 Athletic workbook imported")

    ranks = pd.read_parquet(ranks_path)
    stats = pd.read_parquet(stats_path).set_index("player_name")
    # The workbook's own `Settings` table, in full rather than approximately: every
    # rule it prices at anything other than zero. Attempts, completions, carries and
    # targets really are worth 0 there, so this is his scoring and not a stand-in --
    # which matters, because dropping the interception rule alone is enough to
    # manufacture a quarterback override that he did not make.
    points = (stats["rushingYards"].fillna(0) * 0.1
              + stats["receivingYards"].fillna(0) * 0.1
              + stats["receivingReceptions"].fillna(0) * 0.5
              + stats["passingYards"].fillna(0) * 0.04
              + (stats["rushingTouchdowns"].fillna(0)
                 + stats["receivingTouchdowns"].fillna(0)) * 6
              + stats["passingTouchdowns"].fillna(0) * 4
              - stats["passingInterceptions"].fillna(0) * 2)
    frame = pd.DataFrame({
        "primaryPosition": ranks["position"].to_numpy(),
        "ath_pos_rank": ranks["ath_rank_half"].to_numpy(),
        "ATH_Points": ranks["player_name"].map(points).to_numpy(),
    })
    moved = (sp._attach_athletic_override(frame)
             .assign(big=lambda f: f["ath_override"].abs() >= 5)
             .groupby("primaryPosition")["big"].sum().to_dict())
    assert moved["QB"] == 0, "he has started overriding quarterbacks"
    assert moved["WR"] > moved["RB"] > moved["TE"] >= moved["QB"]
